"""Relatório em .docx com seções selecionáveis."""

import io

from docx import Document


def _text(response) -> str:
    document = Document(io.BytesIO(response.content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def test_sections_catalog(client, signup):
    sections = client.get("/report/sections", headers=signup()).json()
    assert {item["group"] for item in sections} == {"Revisão", "Planejamento", "Condução"}
    assert len(sections) == 16


def test_full_report_contains_protocol_and_conducting_data(client, review):
    review_id, headers = review
    base = f"/reviews/{review_id}"
    client.put(f"{base}/objective", headers=headers, json={"objective": "Mapear XAI em sistemas multiagentes."})
    client.post(f"{base}/questions", headers=headers, json={"question": "Quais técnicas de XAI são usadas?"})
    client.post(f"{base}/criteria", headers=headers, json={"criteria_type": "E", "description": "Agente isolado"})
    scopus = next(item for item in client.get(f"{base}/sources/suggested", headers=headers).json() if item["name"] == "Scopus")
    client.post(f"{base}/sources", headers=headers, json={"source_id": scopus["id"]})
    client.post(f"{base}/sources/{scopus['id']}/import", headers=headers,
                files={"file": ("s.bib", b"@article{a, title={Alpha study}, year={2024}}", "application/octet-stream")})
    article_id = client.get(f"{base}/articles", headers=headers).json()["items"][0]["id"]
    client.patch(f"{base}/articles/{article_id}", headers=headers, json={"status": "A"})

    response = client.get(f"{base}/report.docx", headers=headers)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats")
    text = _text(response)
    for expected in ("Revisão de teste", "Mapear XAI em sistemas multiagentes.", "RQ1: Quais técnicas de XAI são usadas?",
                     "Agente isolado", "Scopus", "Alpha study"):
        assert expected in text


def test_selected_sections_only(client, review):
    review_id, headers = review
    client.put(f"/reviews/{review_id}/objective", headers=headers, json={"objective": "Objetivo visível"})
    client.put(f"/reviews/{review_id}/picoc", headers=headers, json={"population": "População oculta"})
    text = _text(client.get(f"/reviews/{review_id}/report.docx?sections=objectives", headers=headers))
    assert "Objetivo visível" in text
    assert "População oculta" not in text
    assert client.get(f"/reviews/{review_id}/report.docx?sections=nope", headers=headers).status_code == 422
