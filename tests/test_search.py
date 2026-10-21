from __future__ import annotations

import io
import json
import zipfile
from collections.abc import Callable

import httpx
import pytest
from docx import Document
from sqlalchemy import select

from rsl.models import Article, SearchRun, UserApiKey
from rsl.search import CONNECTORS, SearchError
from rsl.search.base import Credentials, Filters
from rsl.search.runner import mark_interrupted

ARXIV_FEED = """<?xml version="1.0" encoding="UTF-8"?>
<feed xmlns="http://www.w3.org/2005/Atom" xmlns:opensearch="http://a9.com/-/spec/opensearch/1.1/"
      xmlns:arxiv="http://arxiv.org/schemas/atom">
  <opensearch:totalResults>2</opensearch:totalResults>
  <entry>
    <id>http://arxiv.org/abs/2401.00001v2</id>
    <published>2024-01-02T00:00:00Z</published>
    <title>Large Language Models for
      Systematic Reviews</title>
    <summary>We study screening with LLMs.</summary>
    <author><name>Ana Souza</name></author>
    <author><name>Bruno Lima</name></author>
    <link href="https://arxiv.org/abs/2401.00001v2" rel="alternate" type="text/html"/>
    <category term="cs.SE"/><category term="cs.CL"/>
  </entry>
  <entry>
    <id>http://arxiv.org/abs/2402.00002v1</id>
    <published>2023-05-01T00:00:00Z</published>
    <title>Automating Literature Search</title>
    <summary>A tool.</summary>
    <author><name>Carla Dias</name></author>
    <arxiv:doi>10.1145/1234567</arxiv:doi>
  </entry>
</feed>"""


def openalex_work(number: int) -> dict:
    return {
        "id": f"https://openalex.org/W{number}", "display_name": f"Work {number}", "publication_year": 2022,
        "doi": f"https://doi.org/10.1145/{number}", "type": "article", "type_crossref": "proceedings-article",
        "abstract_inverted_index": {"Testing": [0], "matters": [1, 3], "really": [2]},
        "authorships": [{"author": {"display_name": "Ana Souza"}, "institutions": [{"display_name": "UFX"}]}],
        "primary_location": {"landing_page_url": f"https://dl.acm.org/doi/10.1145/{number}",
                             "source": {"display_name": "ICSE", "host_organization_name": "ACM", "issn_l": "0000-0001"}},
        "biblio": {"volume": "1", "first_page": "10", "last_page": "20"},
        "keywords": [{"display_name": "Software testing"}],
    }


@pytest.fixture
def api(client) -> Callable[[Callable[[httpx.Request], httpx.Response]], list[httpx.Request]]:
    """Route the connectors' HTTP calls to a handler; returns the list of requests seen."""

    def _install(handler: Callable[[httpx.Request], httpx.Response]) -> list[httpx.Request]:
        seen: list[httpx.Request] = []

        def _record(request: httpx.Request) -> httpx.Response:
            seen.append(request)
            return handler(request)

        client.app.state.search_inline = True
        client.app.state.search_transport = httpx.MockTransport(_record)
        client.app.state.search_sleep = lambda _seconds: None
        return seen

    return _install


def add_source(client, review_id: int, headers: dict, name: str, query: str = 'all:"systematic review"') -> int:
    suggested = client.get(f"/reviews/{review_id}/sources/suggested", headers=headers).json()
    source = next(item for item in suggested if item["name"] == name)
    client.post(f"/reviews/{review_id}/sources", headers=headers, json={"source_id": source["id"]})
    client.put(f"/reviews/{review_id}/search/{source['id']}", headers=headers, json={"search_string": query})
    return source["id"]


def run(client, review_id: int, source_id: int, headers: dict, **payload):
    return client.post(f"/reviews/{review_id}/sources/{source_id}/search-runs", headers=headers,
                       json={"mode": "import", **payload})


def test_suggested_sources_carry_their_connector(client, review):
    review_id, headers = review
    connectors = {item["name"]: item["connector"] for item in
                  client.get(f"/reviews/{review_id}/sources/suggested", headers=headers).json()}
    assert connectors["IEEE Xplore"] == "ieee" and connectors["ACM Digital Library"] == "acm_openalex"
    assert connectors["Google Scholar"] is None
    listed = {item["key"]: item for item in client.get("/search/connectors", headers=headers).json()}
    assert listed["ieee"]["key_required"] and not listed["arxiv"]["key_required"]
    assert "OpenAlex" in listed["acm_openalex"]["notice"]


def test_arxiv_preview_imports_nothing_and_import_records_the_run(client, review, api):
    review_id, headers = review
    requests = api(lambda request: httpx.Response(200, text=ARXIV_FEED))
    source_id = add_source(client, review_id, headers, "arXiv")

    preview = client.post(f"/reviews/{review_id}/sources/{source_id}/search-runs", headers=headers,
                          json={"mode": "preview", "year_from": 2020, "year_to": 2024}).json()
    assert preview["status"] == "done" and preview["total_reported"] == 2
    assert [item["title"] for item in preview["preview"]] == ["Large Language Models for Systematic Reviews",
                                                              "Automating Literature Search"]
    assert "submittedDate:[202001010000 TO 202412312359]" in requests[0].url.params["search_query"]
    assert client.get(f"/reviews/{review_id}/articles", headers=headers).json()["total"] == 0

    imported = run(client, review_id, source_id, headers).json()
    assert imported["status"] == "done" and imported["imported"] == 2 and imported["search_result_id"]
    with client.app.state.database.sessions() as db:
        articles = {item.title: item for item in db.scalars(select(Article).where(Article.review_id == review_id))}
    first = articles["Large Language Models for Systematic Reviews"]
    assert first.author == "Ana Souza; Bruno Lima" and first.year == "2024" and first.status == "U"
    assert first.doi == "10.48550/arxiv.2401.00001" and first.keywords == "cs.SE; cs.CL"
    assert articles["Automating Literature Search"].doi == "10.1145/1234567"
    files = client.get(f"/reviews/{review_id}/imports", headers=headers).json()
    arxiv = next(item for item in files if item["source"]["id"] == source_id)
    assert arxiv["articles"] == 2 and arxiv["files"][0]["from_api"] is True


def test_rerun_skips_studies_already_imported_from_the_source(client, review, api):
    review_id, headers = review
    api(lambda request: httpx.Response(200, text=ARXIV_FEED))
    source_id = add_source(client, review_id, headers, "arXiv")
    assert run(client, review_id, source_id, headers).json()["imported"] == 2
    again = run(client, review_id, source_id, headers).json()
    assert again["imported"] == 0 and again["skipped_existing"] == 2
    assert client.get(f"/reviews/{review_id}/articles", headers=headers).json()["total"] == 2


def test_acm_via_openalex_pages_with_cursor_and_respects_the_limit(client, review, api):
    review_id, headers = review

    def handler(request: httpx.Request) -> httpx.Response:
        cursor = request.url.params["cursor"]
        if cursor == "*":
            return httpx.Response(200, json={"meta": {"count": 5, "next_cursor": "page2"},
                                             "results": [openalex_work(1), openalex_work(2)]})
        return httpx.Response(200, json={"meta": {"count": 5, "next_cursor": "page3"},
                                         "results": [openalex_work(3), openalex_work(4)]})

    requests = api(handler)
    source_id = add_source(client, review_id, headers, "ACM Digital Library", '"software testing", LLM')
    result = run(client, review_id, source_id, headers, max_records=3, year_from=2020).json()
    assert result["status"] == "done" and result["total_reported"] == 5 and result["imported"] == 3
    assert len(requests) == 2
    filters = requests[0].url.params["filter"]
    assert "title_and_abstract.search:\"software testing\"  LLM" in filters
    assert "publication_year:2020-2100" in filters and "publisher_lineage:P4310319798" in filters
    with client.app.state.database.sessions() as db:
        article = db.scalar(select(Article).where(Article.review_id == review_id, Article.title == "Work 1"))
    assert article.abstract == "Testing matters really matters"
    assert article.pages == "10-20" and article.journal == "ICSE" and article.document_type == "inproceedings"
    assert article.doi == "10.1145/1"


def test_failure_imports_nothing_and_shows_the_reason(client, review, api):
    review_id, headers = review
    api(lambda request: httpx.Response(400, text="Bad query syntax"))
    source_id = add_source(client, review_id, headers, "ACM Digital Library")
    result = run(client, review_id, source_id, headers).json()
    assert result["status"] == "failed" and "sintaxe" in result["error"]
    assert client.get(f"/reviews/{review_id}/articles", headers=headers).json()["total"] == 0
    assert client.get(f"/reviews/{review_id}/imports", headers=headers).json()[0]["files"] == []


def test_rate_limit_is_retried_before_giving_up(client, review, api):
    review_id, headers = review
    answers = iter([httpx.Response(429), httpx.Response(200, text=ARXIV_FEED)])
    requests = api(lambda request: next(answers))
    source_id = add_source(client, review_id, headers, "arXiv")
    assert run(client, review_id, source_id, headers).json()["imported"] == 2
    assert len(requests) == 2

    api(lambda request: httpx.Response(429))
    failed = run(client, review_id, source_id, headers).json()
    assert failed["status"] == "failed" and "cota" in failed["error"]


def test_api_keys_are_encrypted_never_returned_and_used_by_their_owner(client, review, api, signup):
    review_id, headers = review
    requests = api(lambda request: httpx.Response(200, json={"total_records": 1, "articles": [{
        "article_number": "99", "title": "IEEE paper", "publication_year": 2021, "content_type": "Conferences",
        "authors": {"authors": [{"full_name": "Ana Souza", "affiliation": "UFX"}]}, "doi": "10.1109/X.1",
        "index_terms": {"author_terms": {"terms": ["LLM"]}, "ieee_terms": {"terms": ["Software"]}},
        "start_page": "1", "end_page": "9", "publication_title": "ICSE"}]}))
    source_id = add_source(client, review_id, headers, "IEEE Xplore", '"Abstract":"systematic review"')

    missing = run(client, review_id, source_id, headers)
    assert missing.status_code == 422 and "Perfil" in missing.json()["detail"]

    assert client.put("/me/api-keys/ieee", headers=headers, json={"key": "chave-teste-ieee-1234"}).status_code == 204
    listed = {item["provider"]: item for item in client.get("/me/api-keys", headers=headers).json()}
    assert listed["ieee"]["configured"] and listed["ieee"]["last_four"] == "1234"
    assert "chave-teste-ieee-1234" not in json.dumps(listed)
    with client.app.state.database.sessions() as db:
        stored = db.scalar(select(UserApiKey).where(UserApiKey.provider == "ieee"))
    assert "chave-teste" not in stored.secret

    result = run(client, review_id, source_id, headers, year_from=2019, year_to=2022).json()
    assert result["status"] == "done" and result["imported"] == 1
    assert requests[-1].url.params["apikey"] == "chave-teste-ieee-1234"
    assert requests[-1].url.params["start_year"] == "2019"
    with client.app.state.database.sessions() as db:
        article = db.scalar(select(Article).where(Article.title == "IEEE paper"))
    assert article.document_type == "inproceedings" and article.author_keywords == "LLM" and article.keywords == "Software"

    other = signup("outra")
    assert not {item["provider"]: item for item in client.get("/me/api-keys", headers=other).json()}["ieee"]["configured"]
    assert client.get(f"/reviews/{review_id}/search-runs", headers=other).status_code == 404
    assert client.post(f"/reviews/{review_id}/sources/{source_id}/search-runs", headers=other,
                       json={"mode": "preview"}).status_code == 404

    assert client.delete("/me/api-keys/ieee", headers=headers).status_code == 204
    assert not {item["provider"]: item for item in client.get("/me/api-keys", headers=headers).json()}["ieee"]["configured"]


def test_key_test_reports_a_rejected_key(client, review, api):
    _review_id, headers = review
    api(lambda request: httpx.Response(401))
    client.put("/me/api-keys/springer", headers=headers, json={"key": "chave-invalida"})
    response = client.post("/me/api-keys/springer/test", headers=headers)
    assert response.status_code == 422 and "recusou a chave" in response.json()["detail"]


def test_run_requires_a_saved_string_and_a_connector(client, review, api):
    review_id, headers = review
    api(lambda request: httpx.Response(200, text=ARXIV_FEED))
    source_id = add_source(client, review_id, headers, "arXiv", "")
    assert run(client, review_id, source_id, headers).status_code == 422
    scholar = add_source(client, review_id, headers, "Google Scholar", "x")
    assert "não tem busca automática" in run(client, review_id, scholar, headers).json()["detail"]


def test_custom_source_may_use_only_open_aggregators(client, review):
    review_id, headers = review
    created = client.post(f"/reviews/{review_id}/sources", headers=headers,
                          json={"name": "Base própria", "connector": "openalex"})
    assert created.status_code == 201 and created.json()["connector"] == "openalex"
    refused = client.post(f"/reviews/{review_id}/sources", headers=headers, json={"name": "Outra", "connector": "ieee"})
    assert refused.status_code == 422


def test_interrupted_runs_are_marked_failed_on_startup(client, review, api):
    review_id, headers = review
    api(lambda request: httpx.Response(200, text=ARXIV_FEED))
    source_id = add_source(client, review_id, headers, "arXiv")
    with client.app.state.database.sessions() as db:
        db.add(SearchRun(review_id=review_id, source_id=source_id, connector="arxiv", query="x", max_records=10,
                         mode="import", status="running", warnings=[]))
        db.commit()
    assert run(client, review_id, source_id, headers).status_code == 409
    mark_interrupted(client.app.state.database)
    runs = client.get(f"/reviews/{review_id}/search-runs", headers=headers).json()
    assert runs[0]["status"] == "failed" and "reinício" in runs[0]["error"]


# ---- keyed connectors, checked on their own --------------------------------------

def pages(connector_key: str, handler, credentials: Credentials, limit: int = 100, filters: Filters | None = None):
    connector = CONNECTORS[connector_key](sleep=lambda _seconds: None)
    with httpx.Client(transport=httpx.MockTransport(handler)) as http:
        return list(connector.pages(http, "query", filters or Filters(), credentials, limit))


def test_scopus_pages_by_cursor_and_warns_without_institutional_token():
    calls = []

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(request)
        cursor = request.url.params["cursor"]
        entries = [{"eid": f"2-s2.0-{cursor}", "dc:title": f"Scopus {cursor}", "dc:creator": "Souza A.",
                    "prism:publicationName": "IST", "prism:coverDate": "2021-03-01", "prism:doi": f"10.1016/{cursor}",
                    "subtype": "ar", "link": [{"@ref": "scopus", "@href": "https://www.scopus.com/x"}]}]
        return httpx.Response(200, json={"search-results": {
            "opensearch:totalResults": "2", "cursor": {"@next": "B" if cursor == "*" else None}, "entry": entries}})

    result = pages("scopus", handler, Credentials(key="k"), filters=Filters(2018, 2022))
    assert [page.records[0]["title"] for page in result] == ["Scopus *", "Scopus B"]
    assert calls[0].headers["X-ELS-APIKey"] == "k" and calls[0].url.params["view"] == "STANDARD"
    assert calls[0].url.params["date"] == "2018-2022"
    assert result[0].warnings and result[0].records[0]["document_type"] == "article"

    complete = pages("scopus", lambda request: httpx.Response(200, json={"search-results": {
        "opensearch:totalResults": "0", "entry": [{"error": "Result set was empty"}]}}), Credentials(key="k", extra="inst"))
    assert complete[0].records == [] and complete[0].warnings == []


def test_sciencedirect_springer_and_wos_map_their_fields():
    sciencedirect = pages("sciencedirect", lambda request: httpx.Response(200, json={"resultsFound": 1, "results": [{
        "title": "SD paper", "authors": [{"name": "Ana"}], "publicationName": "JSS", "publicationDate": "2020-01-01",
        "doi": "10.1016/j.jss.1", "uri": "https://www.sciencedirect.com/x", "pages": {"first": "1", "last": "5"},
        "pii": "S01"}]}), Credentials(key="k"))
    assert sciencedirect[0].records[0]["pages"] == "1-5" and sciencedirect[0].warnings

    def springer(request: httpx.Request) -> httpx.Response:
        assert "onlinedatefrom:2019-01-01" in request.url.params["q"]
        return httpx.Response(200, json={"result": [{"total": "1"}], "records": [{
            "title": "Springer paper", "creators": [{"creator": "Lima, Bruno"}], "publicationName": "EMSE",
            "publicationDate": "2019-05-01", "doi": "10.1007/x", "abstract": {"h1": "Abstract", "p": "Text."},
            "contentType": "Article", "keyword": ["LLM"], "url": [{"format": "html", "value": "https://link.springer.com/x"}]}]})

    record = pages("springer", springer, Credentials(key="k"), filters=Filters(2019, None))[0].records[0]
    assert record["abstract"] == "Text." and record["url"] == "https://link.springer.com/x" and record["keywords"] == "LLM"

    def wos(request: httpx.Request) -> httpx.Response:
        assert request.headers["X-ApiKey"] == "k" and request.url.params["q"].endswith("AND PY=(2015-2100)")
        return httpx.Response(200, json={"metadata": {"total": 1}, "hits": [{
            "uid": "WOS:0001", "title": "WoS paper", "types": ["Proceedings Paper"],
            "source": {"sourceTitle": "ICSE", "publishYear": 2018, "pages": {"range": "3-4"}},
            "names": {"authors": [{"displayName": "Dias, Carla"}]}, "identifiers": {"doi": "10.1/W"},
            "keywords": {"authorKeywords": ["testing"]}, "links": {"record": "https://www.webofscience.com/x"}}]})

    record = pages("wos", wos, Credentials(key="k"), filters=Filters(2015, None))[0].records[0]
    assert record["document_type"] == "inproceedings" and record["author_keywords"] == "testing"
    assert record["bibtex_key"] == "WOS_0001"


def test_rejected_key_is_reported_in_portuguese():
    with pytest.raises(SearchError, match="recusou a chave"):
        pages("ieee", lambda request: httpx.Response(403), Credentials(key="k"))


def test_report_lists_the_executed_search_strategy(client, review, api):
    review_id, headers = review
    api(lambda request: httpx.Response(200, json={"meta": {"count": 1, "next_cursor": None}, "results": [openalex_work(7)]}))
    acm = add_source(client, review_id, headers, "ACM Digital Library", '"software testing" AND LLM')
    assert run(client, review_id, acm, headers, year_from=2020).json()["imported"] == 1
    scopus = add_source(client, review_id, headers, "Scopus", 'TITLE-ABS-KEY("software testing")')
    bib = b"@article{a1, title={Manual study}, author={Lima, Bruno}, year={2021}}"
    client.post(f"/reviews/{review_id}/sources/{scopus}/import", headers=headers, files={"file": ("scopus.bib", bib)})

    docx = client.post(f"/reviews/{review_id}/prisma/export", headers=headers, json={"format": "docx"}).content
    cells = [cell.text for table in Document(io.BytesIO(docx)).tables for row in table.rows for cell in row.cells]
    assert "API (ACM Digital Library (via OpenAlex))" in cells and '"software testing" AND LLM' in cells
    assert "2020–…" in cells and "Arquivo importado (scopus.bib)" in cells
    paragraphs = [paragraph.text for paragraph in Document(io.BytesIO(docx)).paragraphs]
    assert "Observação sobre a ACM Digital Library" in paragraphs

    latex = client.post(f"/reviews/{review_id}/prisma/export", headers=headers, json={"format": "latex"}).content
    main = zipfile.ZipFile(io.BytesIO(latex)).read("main.tex").decode("utf-8")
    assert "Estratégia de busca executada" in main and "TITLE-ABS-KEY" in main
