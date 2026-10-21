"""Relato PRISMA 2020 (spec 010 adaptada): fluxograma, seções, checklist, referências, DOCX/LaTeX e versões."""

import io
import re
import zipfile

import pytest
from docx import Document

BIBTEX = rb"""@article{doe2024, title={Explainable {MAS}}, author={Doe, Jane and Smith, Alex}, journal={J. XAI}, year={2024}, volume={3}, pages={1--9}, doi={10.1/x}}
@inproceedings{roe2023, title={Trust \& agents}, author={Roe, Richard}, booktitle={AAMAS}, year={2023}}
@article{doe2024, title={Same key other paper}, author={Lee, Kim}, journal={AI}, year={2022}}
@article{dup, title={Explainable MAS}, author={Doe, Jane}, year={2024}}
@article{rej, title={Off topic}, author={X, Y}, year={2020}}
@article{pend, title={Pending}, author={Z, W}, year={2021}}
"""


@pytest.fixture
def report(client, review):
    review_id, headers = review
    base = f"/reviews/{review_id}"
    scopus = next(item for item in client.get(f"{base}/sources/suggested", headers=headers).json() if item["name"] == "Scopus")
    client.post(f"{base}/sources", headers=headers, json={"source_id": scopus["id"]})
    client.post(f"{base}/sources/{scopus['id']}/import", headers=headers, files={"file": ("s.bib", BIBTEX, "application/octet-stream")})
    items = client.get(f"{base}/articles?page_size=50", headers=headers).json()["items"]
    by_key = {(item["bibtex_key"], item["title"]): item["id"] for item in items}
    ids = {title: article_id for (_, title), article_id in by_key.items()}
    criteria = client.post(f"{base}/criteria", headers=headers, json={"criteria_type": "E", "description": "Fora do tema"}).json()
    main = by_key[("doe2024", "Explainable MAS")]
    trust = next(article_id for (_, title), article_id in by_key.items() if title.startswith("Trust"))
    client.post(f"{base}/articles/bulk", headers=headers, json={"ids": [main, trust, ids["Same key other paper"]], "action": "accept"})
    client.patch(f"{base}/articles/{ids['Off topic']}", headers=headers, json={"status": "R", "selection_criteria_id": criteria["id"]})
    client.patch(f"{base}/articles/{by_key[('dup', 'Explainable MAS')]}", headers=headers, json={"status": "D"})
    # Quality: "Same key other paper" falls below the cutoff.
    question = client.post(f"{base}/quality-questions", headers=headers, json={"description": "Claro?"}).json()
    answers = {item["description"]: item["id"] for item in client.post(f"{base}/quality-answers/suggested", headers=headers).json()}
    client.put(f"{base}/quality-cutoff", headers=headers, json={"quality_cutoff_score": 0.5})
    for article_id, label in ((main, "Sim"), (trust, "Parcialmente"), (ids["Same key other paper"], "Não")):
        client.put(f"{base}/articles/{article_id}/quality", headers=headers, json={"question_id": question["id"], "answer_id": answers[label]})
    return base, headers


def test_flow_numbers_add_up(client, report):
    base, headers = report
    flow = client.get(f"{base}/prisma", headers=headers).json()["flow"]
    assert (flow["identified"], flow["duplicates_removed"], flow["screened"]) == (6, 1, 5)
    assert (flow["excluded_screening"], flow["awaiting_screening"]) == (1, 1)
    assert flow["excluded_by_criteria"] == [{"criteria": "Fora do tema", "count": 1}]
    assert (flow["assessed_quality"], flow["excluded_quality"], flow["included"]) == (3, 1, 2)
    assert flow["consistent"] is True


def test_sections_required_warnings_and_checklist_rules(client, report):
    base, headers = report
    status = client.get(f"{base}/prisma", headers=headers).json()
    assert any("Ameaças à validade" in warning for warning in status["warnings"])
    assert len(status["checklist"]) == 42
    assert client.put(f"{base}/prisma/sections/nope", headers=headers, json={"content": "x"}).status_code == 404
    for key in ("threats", "ai_statement"):
        client.put(f"{base}/prisma/sections/{key}", headers=headers, json={"content": "Texto dos pesquisadores."})
    assert not any("obrigatória" in warning for warning in client.get(f"{base}/prisma", headers=headers).json()["warnings"])

    url = f"{base}/prisma/checklist"
    assert client.put(f"{url}/1", headers=headers, json={"status": "met"}).status_code == 422
    assert client.put(f"{url}/12", headers=headers, json={"status": "na"}).status_code == 422
    assert client.put(f"{url}/99", headers=headers, json={"status": "na", "justification": "x"}).status_code == 404
    met = client.put(f"{url}/1", headers=headers, json={"status": "met", "location": "Título"}).json()
    assert (met["status_label"], met["location"]) == ("Atendido", "Título")
    assert client.put(f"{url}/12", headers=headers, json={"status": "na", "justification": "Síntese qualitativa"}).json()["status"] == "na"


def test_references_in_both_formats(client, report):
    base, headers = report
    bib = client.get(f"{base}/references?format=bibtex", headers=headers).text
    assert bib.count("@") == 2 and "@inproceedings{roe2023" in bib and "@article{doe2024" in bib
    ris = client.get(f"{base}/references?format=ris", headers=headers).text
    assert ris.count("TY  - ") == 2 and "AU  - Doe, Jane" in ris


def _docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    return "\n".join([p.text for p in document.paragraphs] + [c.text for t in document.tables for r in t.rows for c in r.cells])


def test_docx_export_has_skeleton_flow_and_appendix_and_creates_version(client, report):
    base, headers = report
    client.put(f"{base}/prisma/sections/introduction", headers=headers, json={"content": "Primeiro parágrafo.\n\nSegundo parágrafo."})
    response = client.post(f"{base}/prisma/export", headers=headers, json={"format": "docx"})
    assert response.status_code == 200
    text = _docx_text(response.content)
    for expected in ("Introdução", "Segundo parágrafo.", "Ameaças à validade", "[A ser escrito pelos pesquisadores.]",
                     "Registros identificados (n = 6)", "Estudos incluídos na revisão (n = 2)", "Apêndice — Checklist PRISMA 2020",
                     "Doe, Jane, Smith, Alex (2024). Explainable MAS. J. XAI, 3, 1--9. https://doi.org/10.1/x"):
        assert expected in text
    versions = client.get(f"{base}/prisma/versions", headers=headers).json()
    assert len(versions) == 1 and versions[0]["numbers"]["included"] == 2
    download = client.get(f"{base}/prisma/versions/{versions[0]['id']}/download", headers=headers)
    assert download.content == response.content


def test_abnt_style_is_used_in_docx_and_latex(client, report):
    base, headers = report
    client.put(f"{base}/prisma/style", headers=headers, json={"reference_style": "abnt"})
    text = _docx_text(client.post(f"{base}/prisma/export", headers=headers, json={"format": "docx"}).content)
    assert "DOE, Jane; SMITH, Alex. Explainable MAS. J. XAI, v. 3, p. 1--9, 2024." in text
    archive = zipfile.ZipFile(io.BytesIO(client.post(f"{base}/prisma/export", headers=headers, json={"format": "latex"}).content))
    assert r"\bibliographystyle{abntex2-alf}" in archive.read("main.tex").decode()


def test_latex_export_keeps_citation_keys_and_bib(client, report):
    base, headers = report
    response = client.post(f"{base}/prisma/export", headers=headers, json={"format": "latex"})
    archive = zipfile.ZipFile(io.BytesIO(response.content))
    assert sorted(archive.namelist()) == ["main.tex", "referencias.bib"]
    main = archive.read("main.tex").decode()
    bib = archive.read("referencias.bib").decode()
    assert r"\begin{tikzpicture}" in main and r"\bibliography{referencias}" in main
    cited = {key for group in re.findall(r"\\cite\{([^}]*)\}", main) for key in group.split(",")}
    declared = set(re.findall(r"@\w+\{([^,]+),", bib))
    assert cited and cited <= declared
    assert r"Trust \& agents" in main and r"\textbackslash" not in main
    assert r"\bibliographystyle{apalike}" in main
