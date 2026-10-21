"""Reporting: Word (.docx) export with selectable sections, as Parsifal's Export page."""

from __future__ import annotations

import io
from collections.abc import Callable
from datetime import datetime, timezone

from docx import Document
from docx.document import Document as DocxDocument
from fastapi import APIRouter, HTTPException, Response
from sqlalchemy import select

from ..deps import DbSession, MemberReview
from ..models import ARTICLE_STATUSES, Article, Review, SearchSession
from .analysis import analysis_data, extraction_rows, quality_scores
from .conducting import import_summary
from .planning import planning_out

router = APIRouter(tags=["relatório"])

RELATED_TO = {"": "—", "P": "População", "I": "Intervenção", "C": "Comparação", "O": "Resultado"}


def _table(doc: DocxDocument, header: list[str], rows: list[list[str]]) -> None:
    if not rows:
        doc.add_paragraph("Nenhum registro.")
        return
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = text
    for row in rows:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = str(text)


def _bullets(doc: DocxDocument, items: list[str]) -> None:
    if not items:
        doc.add_paragraph("Nenhum registro.")
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _text(doc: DocxDocument, text: str) -> None:
    doc.add_paragraph(text or "Não informado.")


def section_review(doc: DocxDocument, db: DbSession, review: Review) -> None:
    doc.add_heading("Revisão", level=1)
    _text(doc, review.description)
    authors = [f"{review.author.full_name} (autor principal)", *(user.full_name for user in review.co_authors)]
    doc.add_heading("Autores", level=2)
    _bullets(doc, authors)


def section_objectives(doc: DocxDocument, db: DbSession, review: Review) -> None:
    doc.add_heading("Objetivos", level=2)
    _text(doc, review.objective)


def section_picoc(doc: DocxDocument, db: DbSession, review: Review) -> None:
    doc.add_heading("PICOC", level=2)
    _table(doc, ["Elemento", "Descrição"], [
        ["População", review.population], ["Intervenção", review.intervention], ["Comparação", review.comparison],
        ["Resultado", review.outcome], ["Contexto", review.context],
    ])


def section_questions(doc: DocxDocument, db: DbSession, review: Review) -> None:
    doc.add_heading("Perguntas de pesquisa", level=2)
    questions = planning_out(db, review)["questions"]
    _bullets(doc, [f"RQ{index}: {item['question']}" for index, item in enumerate(questions, start=1)])


def section_keywords(doc: DocxDocument, db: DbSession, review: Review) -> None:
    doc.add_heading("Palavras-chave e sinônimos", level=2)
    _table(doc, ["Palavra-chave", "Sinônimos", "Relacionada a"], [
        [item["description"], ", ".join(synonym["description"] for synonym in item["synonyms"]),
         RELATED_TO[item["related_to"]]]
        for item in planning_out(db, review)["keywords"]
    ])


def section_search_string(doc: DocxDocument, db: DbSession, review: Review) -> None:
    doc.add_heading("String de busca", level=2)
    _text(doc, planning_out(db, review)["search_string"])


def section_sources(doc: DocxDocument, db: DbSession, review: Review) -> None:
    doc.add_heading("Fontes", level=2)
    _table(doc, ["Fonte", "URL"], [[source.name, source.url] for source in review.sources])


def section_criteria(doc: DocxDocument, db: DbSession, review: Review) -> None:
    planning = planning_out(db, review)
    doc.add_heading("Critérios de inclusão", level=2)
    _bullets(doc, [item["description"] for item in planning["inclusion_criteria"]])
    doc.add_heading("Critérios de exclusão", level=2)
    _bullets(doc, [item["description"] for item in planning["exclusion_criteria"]])


def section_quality_checklist(doc: DocxDocument, db: DbSession, review: Review) -> None:
    planning = planning_out(db, review)
    doc.add_heading("Checklist de qualidade", level=2)
    _bullets(doc, [f"Q{index}: {item['description']}" for index, item in enumerate(planning["quality_questions"], start=1)])
    _table(doc, ["Resposta", "Peso"], [[item["description"], f"{item['weight']:g}"] for item in planning["quality_answers"]])
    doc.add_paragraph(f"Pontuação máxima: {planning['quality_max_score']:g} · Nota de corte: {planning['quality_cutoff_score']:g}")


def section_extraction_form(doc: DocxDocument, db: DbSession, review: Review) -> None:
    doc.add_heading("Formulário de extração", level=2)
    _table(doc, ["Campo", "Tipo", "Valores"], [
        [item["description"], item["field_type_label"], ", ".join(lookup["value"] for lookup in item["lookups"])]
        for item in planning_out(db, review)["data_extraction_fields"]
    ])


def section_source_strings(doc: DocxDocument, db: DbSession, review: Review) -> None:
    doc.add_heading("Strings por fonte", level=2)
    strings = {
        session.source_id: session.search_string
        for session in db.scalars(select(SearchSession).where(SearchSession.review_id == review.id))
    }
    _table(doc, ["Fonte", "String"], [[source.name, strings.get(source.id, "")] for source in review.sources])


def section_imports(doc: DocxDocument, db: DbSession, review: Review) -> None:
    doc.add_heading("Estudos importados", level=2)
    _table(doc, ["Fonte", "Estudos", "Arquivos"], [
        [item["source"]["name"], str(item["articles"]), ", ".join(file["file_name"] for file in item["files"])]
        for item in import_summary(review, db)
    ])


def section_selection(doc: DocxDocument, db: DbSession, review: Review) -> None:
    data = analysis_data(db, review)
    doc.add_heading("Seleção de estudos", level=2)
    _table(doc, ["Status", "Estudos"], [[ARTICLE_STATUSES[code], str(count)] for code, count in data["status_totals"].items()])
    _table(doc, ["Fonte", "Estudos", "Aceitos"],
           [[item["source"], str(item["total"]), str(item["accepted"])] for item in data["per_source"]])
    doc.add_heading("Estudos aceitos", level=3)
    accepted = db.scalars(
        select(Article).where(Article.review_id == review.id, Article.status == "A").order_by(Article.title)
    ).all()
    _bullets(doc, [
        f"{article.author or 'Autor não informado'} ({article.year or 's.d.'}). {article.title}."
        + (f" {article.journal}." if article.journal else "")
        for article in accepted
    ])


def section_quality_assessment(doc: DocxDocument, db: DbSession, review: Review) -> None:
    doc.add_heading("Avaliação de qualidade", level=2)
    scores = quality_scores(db, review)
    articles = {article.id: article for article in db.scalars(select(Article).where(Article.id.in_(scores)))} if scores else {}
    _table(doc, ["Estudo", "Pontuação", "Acima do corte"], [
        [articles[article_id].title, f"{entry['score']:g}", "Sim" if entry["score"] >= review.quality_cutoff_score else "Não"]
        for article_id, entry in sorted(scores.items(), key=lambda item: articles[item[0]].title)
    ])


def section_data_extraction(doc: DocxDocument, db: DbSession, review: Review) -> None:
    doc.add_heading("Extração de dados", level=2)
    header, rows = extraction_rows(db, review)
    _table(doc, header, rows)


def section_analysis(doc: DocxDocument, db: DbSession, review: Review) -> None:
    data = analysis_data(db, review)
    doc.add_heading("Análise de dados", level=2)
    doc.add_paragraph(f"Seleção final: {data['final_total']} estudo(s).")
    _table(doc, ["Ano", "Estudos"], [[item["year"], str(item["count"])] for item in data["final_per_year"]])


Section = tuple[str, str, str, Callable[[DocxDocument, DbSession, Review], None]]
SECTIONS: list[Section] = [
    ("review", "Revisão", "Revisão (descrição e autores)", section_review),
    ("objectives", "Planejamento", "Objetivos", section_objectives),
    ("picoc", "Planejamento", "PICOC", section_picoc),
    ("questions", "Planejamento", "Perguntas de pesquisa", section_questions),
    ("keywords", "Planejamento", "Palavras-chave e sinônimos", section_keywords),
    ("search_string", "Planejamento", "String de busca", section_search_string),
    ("sources", "Planejamento", "Fontes", section_sources),
    ("criteria", "Planejamento", "Critérios de seleção", section_criteria),
    ("quality_checklist", "Planejamento", "Checklist de qualidade", section_quality_checklist),
    ("extraction_form", "Planejamento", "Formulário de extração", section_extraction_form),
    ("source_strings", "Condução", "Strings por fonte", section_source_strings),
    ("imports", "Condução", "Estudos importados", section_imports),
    ("selection", "Condução", "Seleção de estudos", section_selection),
    ("quality_assessment", "Condução", "Avaliação de qualidade", section_quality_assessment),
    ("data_extraction", "Condução", "Extração de dados", section_data_extraction),
    ("analysis", "Condução", "Análise de dados", section_analysis),
]


@router.get("/report/sections")
def report_sections() -> list[dict[str, str]]:
    return [{"key": key, "group": group, "label": label} for key, group, label, _ in SECTIONS]


@router.get("/reviews/{review_id}/report.docx")
def export_report(review: MemberReview, db: DbSession, sections: str = "") -> Response:
    wanted = [item for item in sections.split(",") if item] or [key for key, *_ in SECTIONS]
    unknown = set(wanted) - {key for key, *_ in SECTIONS}
    if unknown:
        raise HTTPException(status_code=422, detail=f"Seções desconhecidas: {', '.join(sorted(unknown))}")
    doc = Document()
    doc.add_heading(review.title, level=0)
    doc.add_paragraph(f"Gerado em {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M')} UTC")
    current_group = None
    for key, group, _, render in SECTIONS:
        if key not in wanted:
            continue
        if group != current_group and group != "Revisão":
            doc.add_heading(group, level=1)
        current_group = group
        render(doc, db, review)
    buffer = io.BytesIO()
    doc.save(buffer)
    return Response(
        buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{review.name}.docx"'},
    )
