"""PRISMA 2020 report (spec 010 adapted): flow diagram, skeleton, checklist, DOCX/LaTeX export and versions."""

from __future__ import annotations

import io
import re
import secrets
import unicodedata
import zipfile
from collections import Counter
from datetime import datetime, timezone
from typing import Any, Literal

from docx import Document
from docx.document import Document as DocxDocument
from docx.shared import Pt
from fastapi import APIRouter, HTTPException, Response
from pydantic import BaseModel
from sqlalchemy import select

from ..deps import AppSettings, CurrentUser, DbSession, MemberReview, log_activity
from ..models import (
    REFERENCE_STYLES,
    REPORT_SECTIONS,
    REQUIRED_SECTIONS,
    Article,
    PrismaItem,
    ReportSection,
    ReportSettings,
    ReportVersion,
    Review,
    SearchSession,
    SelectionCriteria,
)
from ..prisma import PRISMA_ATTRIBUTION, PRISMA_ITEMS, PRISMA_KEYS, PRISMA_STATUSES
from .analysis import _questions, extraction_rows, final_articles, quality_scores
from .insights import Filters, compute
from .planning import planning_out
from .synthesis import agreement_data, list_codes, list_themes

router = APIRouter(tags=["relato PRISMA"])

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


# ---- data ------------------------------------------------------------------------------

def prisma_flow(db: DbSession, review: Review) -> dict[str, Any]:
    """PRISMA 2020 flow numbers computed from the review; the boxes always add up."""
    articles = db.scalars(select(Article).where(Article.review_id == review.id)).all()
    per_source = Counter(article.source.name if article.source else "Sem fonte" for article in articles)
    status = Counter(article.status for article in articles)
    criteria = {item.id: item.description for item in db.scalars(
        select(SelectionCriteria).where(SelectionCriteria.review_id == review.id))}
    excluded_by_criteria = Counter(criteria.get(article.selection_criteria_id, "Sem critério informado")
                                   for article in articles if article.status == "R")
    final = final_articles(db, review)
    identified = len(articles)
    screened = identified - status["D"]
    assessed = status["A"]
    excluded_quality = assessed - len(final) if _questions(db, review) else 0
    return {
        "identified": identified,
        "per_source": [{"source": name, "count": count} for name, count in sorted(per_source.items())],
        "duplicates_removed": status["D"],
        "screened": screened,
        "excluded_screening": status["R"],
        "excluded_by_criteria": [{"criteria": name, "count": count} for name, count in excluded_by_criteria.most_common()],
        "awaiting_screening": status["U"],
        "assessed_quality": assessed,
        "excluded_quality": excluded_quality,
        "included": len(final),
        "consistent": identified - status["D"] - status["R"] - status["U"] - excluded_quality == len(final),
    }


def _settings(db: DbSession, review: Review) -> ReportSettings:
    settings = db.get(ReportSettings, review.id)
    if settings is None:
        settings = ReportSettings(review_id=review.id)
        db.add(settings)
        db.flush()
    return settings


def _sections(db: DbSession, review: Review) -> dict[str, str]:
    stored = {item.key: item.content for item in db.scalars(select(ReportSection).where(ReportSection.review_id == review.id))}
    return {key: stored.get(key, "") for key in REPORT_SECTIONS}


def _checklist(db: DbSession, review: Review) -> list[dict[str, Any]]:
    stored = {item.item: item for item in db.scalars(select(PrismaItem).where(PrismaItem.review_id == review.id))}
    return [
        {"item": item, "section": section, "topic": topic, "description": description,
         "status": stored[item].status if item in stored else None,
         "status_label": PRISMA_STATUSES[stored[item].status] if item in stored else "Pendente",
         "location": stored[item].location if item in stored else "",
         "justification": stored[item].justification if item in stored else ""}
        for item, section, topic, description in PRISMA_ITEMS
    ]


def _warnings(sections: dict[str, str], checklist: list[dict[str, Any]], flow: dict[str, Any]) -> list[str]:
    warnings = [f"A seção '{REPORT_SECTIONS[key]}' é obrigatória e está vazia." for key in REQUIRED_SECTIONS if not sections[key].strip()]
    pending = sum(item["status"] is None for item in checklist)
    if pending:
        warnings.append(f"{pending} item(ns) do checklist PRISMA ainda sem avaliação.")
    if flow["awaiting_screening"]:
        warnings.append(f"{flow['awaiting_screening']} estudo(s) ainda sem classificação na seleção.")
    return warnings


# ---- references ---------------------------------------------------------------------------

def _ascii(text: str) -> str:
    return unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")


def citation_keys(articles: list[Article]) -> dict[int, str]:
    keys: dict[int, str] = {}
    used: set[str] = set()
    for article in articles:
        base = re.sub(r"[^A-Za-z0-9:_-]", "", _ascii(article.bibtex_key)) or f"art{article.id}"
        key, suffix = base, 2
        while key in used:
            key, suffix = f"{base}{chr(96 + suffix) if suffix < 27 else suffix}", suffix + 1
        used.add(key)
        keys[article.id] = key
    return keys


def _authors(article: Article) -> list[str]:
    return [part.strip() for part in re.split(r"\s+and\s+|;", article.author) if part.strip()]


def format_reference(article: Article, style: str) -> str:
    authors = _authors(article)
    year = article.year or "s.d."
    doi = f" https://doi.org/{article.doi}" if article.doi else ""
    if style == "abnt":
        def abnt(name: str) -> str:
            if "," in name:
                surname, given = (piece.strip() for piece in name.split(",", 1))
            else:
                parts = name.split()
                surname, given = (parts[-1], " ".join(parts[:-1])) if len(parts) > 1 else (name, "")
            return f"{surname.upper()}, {given}".strip().rstrip(",")
        names = "; ".join(abnt(name) for name in authors) or "AUTOR DESCONHECIDO"
        venue = f" {article.journal}," if article.journal else ""
        volume = f" v. {article.volume}," if article.volume else ""
        pages = f" p. {article.pages}," if article.pages else ""
        return f"{names}. {article.title}.{venue}{volume}{pages} {year}.{(' DOI:' + doi) if doi else ''}".replace("  ", " ")
    names = ", ".join(authors) or "Autor desconhecido"
    venue = f" {article.journal}" if article.journal else ""
    volume = f", {article.volume}" if article.volume else ""
    pages = f", {article.pages}" if article.pages else ""
    return f"{names} ({year}). {article.title}.{venue}{volume}{pages}.{doi}".replace("..", ".")


def bibtex(articles: list[Article], keys: dict[int, str]) -> str:
    entries = []
    for article in articles:
        kind = "inproceedings" if re.sub(r"[^a-z]", "", article.document_type.casefold()) in {"inproceedings", "conf", "cpaper", "conference"} else "article"
        venue_field = "booktitle" if kind == "inproceedings" else "journal"
        fields = {"title": article.title, "author": article.author, venue_field: article.journal, "year": article.year,
                  "volume": article.volume, "pages": article.pages, "doi": article.doi, "url": article.url}
        body = ",\n".join(f"  {name} = {{{value.replace('{', '').replace('}', '')}}}" for name, value in fields.items() if value)
        entries.append(f"@{kind}{{{keys[article.id]},\n{body}\n}}")
    return "\n\n".join(entries) + "\n"


def ris(articles: list[Article]) -> str:
    lines = []
    for article in articles:
        kind = "CONF" if re.sub(r"[^a-z]", "", article.document_type.casefold()) in {"inproceedings", "conf", "cpaper"} else "JOUR"
        lines.append(f"TY  - {kind}")
        lines.append(f"TI  - {article.title}")
        lines += [f"AU  - {name}" for name in _authors(article)]
        for tag, value in (("PY", article.year), ("JO", article.journal), ("VL", article.volume), ("SP", article.pages),
                           ("DO", article.doi), ("UR", article.url), ("AB", article.abstract)):
            if value:
                lines.append(f"{tag}  - {value}")
        lines.append("ER  - ")
    return "\n".join(lines) + "\n"


# ---- API: status, sections, checklist, references --------------------------------------------

@router.get("/reviews/{review_id}/prisma")
def get_prisma(review: MemberReview, db: DbSession) -> dict[str, Any]:
    flow = prisma_flow(db, review)
    sections = _sections(db, review)
    checklist = _checklist(db, review)
    style = _settings(db, review).reference_style
    db.commit()
    return {"flow": flow, "sections": sections, "section_labels": REPORT_SECTIONS, "required": list(REQUIRED_SECTIONS),
            "reference_style": style, "reference_styles": REFERENCE_STYLES, "checklist": checklist,
            "statuses": PRISMA_STATUSES, "attribution": PRISMA_ATTRIBUTION,
            "warnings": _warnings(sections, checklist, flow)}


class SectionIn(BaseModel):
    content: str


@router.put("/reviews/{review_id}/prisma/sections/{key}")
def save_section(key: str, payload: SectionIn, review: MemberReview, db: DbSession, user: CurrentUser) -> dict[str, str]:
    if key not in REPORT_SECTIONS:
        raise HTTPException(status_code=404, detail="Seção desconhecida.")
    section = db.get(ReportSection, (review.id, key))
    if section is None:
        section = ReportSection(review_id=review.id, key=key)
        db.add(section)
    section.content, section.updated_by_id = payload.content.strip(), user.id
    db.commit()
    return {"key": key, "content": section.content}


class StyleIn(BaseModel):
    reference_style: Literal["apa", "abnt"]


@router.put("/reviews/{review_id}/prisma/style")
def save_style(payload: StyleIn, review: MemberReview, db: DbSession) -> dict[str, str]:
    _settings(db, review).reference_style = payload.reference_style
    db.commit()
    return {"reference_style": payload.reference_style}


class ChecklistIn(BaseModel):
    status: Literal["met", "not_met", "na"]
    location: str = ""
    justification: str = ""


@router.put("/reviews/{review_id}/prisma/checklist/{item}")
def save_checklist_item(item: str, payload: ChecklistIn, review: MemberReview, db: DbSession, user: CurrentUser) -> dict[str, Any]:
    if item not in PRISMA_KEYS:
        raise HTTPException(status_code=404, detail="Item do PRISMA desconhecido.")
    if payload.status == "met" and not payload.location.strip():
        raise HTTPException(status_code=422, detail="Informe onde o item é atendido no relato (seção ou página).")
    if payload.status != "met" and not payload.justification.strip():
        raise HTTPException(status_code=422, detail="Justifique por que o item não é atendido ou não se aplica.")
    row = db.get(PrismaItem, (review.id, item))
    if row is None:
        row = PrismaItem(review_id=review.id, item=item, status=payload.status)
        db.add(row)
    row.status, row.location, row.justification, row.user_id = (
        payload.status, payload.location.strip(), payload.justification.strip(), user.id)
    db.commit()
    return next(entry for entry in _checklist(db, review) if entry["item"] == item)


@router.get("/reviews/{review_id}/references")
def export_references(review: MemberReview, db: DbSession, format: Literal["bibtex", "ris"] = "bibtex") -> Response:
    articles = final_articles(db, review)
    if format == "ris":
        return Response(ris(articles), media_type="application/x-research-info-systems",
                        headers={"Content-Disposition": f'attachment; filename="{review.name}-referencias.ris"'})
    return Response(bibtex(articles, citation_keys(articles)), media_type="application/x-bibtex",
                    headers={"Content-Disposition": f'attachment; filename="{review.name}-referencias.bib"'})


# ---- report content shared by DOCX and LaTeX ---------------------------------------------------

def report_data(db: DbSession, review: Review) -> dict[str, Any]:
    articles = final_articles(db, review)
    keys = citation_keys(articles)
    scores = quality_scores(db, review)
    style = _settings(db, review).reference_style
    planning = planning_out(db, review)
    strings = {item.source_id: item.search_string for item in db.scalars(select(SearchSession).where(SearchSession.review_id == review.id))}
    themes = list_themes(review, db)
    header, rows = extraction_rows(db, review)
    insights = compute(db, review, Filters())
    agreement = agreement_data(db, review)
    return {
        "review": review, "articles": articles, "keys": keys, "scores": scores, "style": style, "planning": planning,
        "source_strings": [(source.name, strings.get(source.id, "")) for source in review.sources],
        "base_string": strings.get(None, ""), "themes": themes, "codes": [code for code in list_codes(review, db) if code["status"] == "active"],
        "extraction": (header, rows), "hypotheses": insights["hypotheses"], "agreement": agreement,
        "flow": prisma_flow(db, review), "sections": _sections(db, review), "checklist": _checklist(db, review),
        "references": [(keys[article.id], format_reference(article, style)) for article in sorted(articles, key=lambda a: (a.author.casefold(), a.year))],
    }


def methods_blocks(data: dict[str, Any]) -> list[tuple[str, list[str] | list[list[str]], str]]:
    """(title, content, kind) blocks with protocol facts; kind is 'list', 'table' or 'text'."""
    planning = data["planning"]
    picoc = planning["picoc"]
    blocks: list[tuple[str, Any, str]] = [
        ("Objetivos", [planning["objective"] or "Não informado."], "text"),
        ("PICOC", [["População", picoc["population"]], ["Intervenção", picoc["intervention"]], ["Comparação", picoc["comparison"]],
                   ["Resultado", picoc["outcome"]], ["Contexto", picoc["context"]]], "table"),
        ("Perguntas de pesquisa", [f"RQ{index}: {item['question']}" for index, item in enumerate(planning["questions"], 1)], "list"),
        ("Critérios de inclusão", [item["description"] for item in planning["inclusion_criteria"]], "list"),
        ("Critérios de exclusão", [item["description"] for item in planning["exclusion_criteria"]], "list"),
        ("Fontes e strings de busca", [[name, string or data["base_string"] or "—"] for name, string in data["source_strings"]], "table"),
        ("Avaliação de qualidade", [f"Q{index}: {item['description']}" for index, item in enumerate(planning["quality_questions"], 1)]
         + [f"Respostas: {', '.join(f'{a['description']} ({a['weight']:g})' for a in planning['quality_answers'])}" if planning["quality_answers"] else "Sem respostas cadastradas.",
            f"Pontuação máxima {planning['quality_max_score']:g}; nota de corte {planning['quality_cutoff_score']:g}."], "list"),
        ("Formulário de extração", [[field["description"], field["field_type_label"]] for field in planning["data_extraction_fields"]], "table"),
    ]
    agreement = data["agreement"]
    coding = [f"{len(data['codes'])} código(s) ativo(s) e {len(data['themes'])} tema(s)."]
    if agreement["enabled"]:
        value = "—" if agreement["value"] is None else f"{agreement['value']:.3f}"
        coding.append(f"Dupla codificação em amostra: {agreement['segments']} trecho(s); {'κ' if agreement['metric'] == 'kappa' else 'α'} = {value}.")
    blocks.append(("Síntese temática", coding, "list"))
    return blocks


def results_blocks(data: dict[str, Any]) -> list[tuple[str, Any, str]]:
    keys = data["keys"]
    max_score_rows = [[keys[article.id], article.title, f"{data['scores'].get(article.id, {}).get('score', 0):g}"] for article in data["articles"]]
    header, rows = data["extraction"]
    blocks: list[tuple[str, Any, str]] = [
        ("Estudos incluídos", [[keys[article.id], article.title, article.year, article.source.name if article.source else ""]
                               for article in data["articles"]], "table"),
        ("Avaliação de qualidade dos incluídos", max_score_rows, "table") if data["scores"] else ("Avaliação de qualidade", ["Sem checklist de qualidade."], "text"),
        ("Dados extraídos", [header, *rows], "header_table") if len(header) > 5 else ("Dados extraídos", ["Sem campos de extração."], "text"),
        ("Temas", [f"{theme['name']}: " + ", ".join(f"{code['name']} ({len(code['evidence'])} trecho(s))" for code in theme["codes"])
                   for theme in data["themes"]], "list"),
    ]
    for hypothesis in data["hypotheses"]:
        blocks.append((f"Hipótese {hypothesis['code']}: {hypothesis['statement']}",
                       [[item["label"], str(item["count"]), f"{item['pct']}%", f"{item['weighted_pct']}%"] for item in hypothesis["verdicts"]],
                       "table"))
    return blocks


# ---- DOCX ---------------------------------------------------------------------------------------

def _docx_table(doc: DocxDocument, rows: list[list[str]], header: list[str] | None = None) -> None:
    if not rows:
        doc.add_paragraph("Nenhum registro.")
        return
    width = len(header or rows[0])
    table = doc.add_table(rows=0, cols=width)
    table.style = "Table Grid"
    for values in ([header] if header else []) + rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = str(value)


def _docx_flow(doc: DocxDocument, flow: dict[str, Any]) -> None:
    sources = "; ".join(f"{item['source']} (n = {item['count']})" for item in flow["per_source"])
    criteria = "; ".join(f"{item['criteria']} (n = {item['count']})" for item in flow["excluded_by_criteria"])
    rows = [
        ["Identificação", f"Registros identificados (n = {flow['identified']})\n{sources}", f"Duplicados removidos (n = {flow['duplicates_removed']})"],
        ["Triagem", f"Registros triados (n = {flow['screened']})",
         f"Excluídos na seleção (n = {flow['excluded_screening']}){chr(10) + criteria if criteria else ''}\nAguardando classificação (n = {flow['awaiting_screening']})"],
        ["Elegibilidade", f"Avaliados na qualidade (n = {flow['assessed_quality']})", f"Excluídos pela nota de corte (n = {flow['excluded_quality']})"],
        ["Inclusão", f"Estudos incluídos na revisão (n = {flow['included']})", ""],
    ]
    _docx_table(doc, rows)


def build_docx(data: dict[str, Any]) -> bytes:
    review, sections = data["review"], data["sections"]
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(sections["title"] or review.title, level=0)
    for key, label in REPORT_SECTIONS.items():
        if key == "title":
            continue
        doc.add_heading(label, level=1)
        text = sections[key]
        if text:
            for paragraph in re.split(r"\n\s*\n", text):
                doc.add_paragraph(paragraph.strip())
        else:
            doc.add_paragraph("[A ser escrito pelos pesquisadores.]").runs[0].italic = True
        if key in {"methods", "results"}:
            if key == "results":
                doc.add_heading("Fluxograma PRISMA 2020", level=2)
                _docx_flow(doc, data["flow"])
            for title, content, kind in (methods_blocks(data) if key == "methods" else results_blocks(data)):
                doc.add_heading(title, level=2)
                if kind == "table":
                    _docx_table(doc, content)
                elif kind == "header_table":
                    _docx_table(doc, content[1:], header=content[0])
                elif kind == "list":
                    if not content:
                        doc.add_paragraph("Nenhum registro.")
                    for item in content:
                        doc.add_paragraph(item, style="List Bullet")
                else:
                    for item in content:
                        doc.add_paragraph(item)
    doc.add_heading("Referências", level=1)
    for _, reference in data["references"]:
        doc.add_paragraph(reference)
    doc.add_heading("Apêndice — Checklist PRISMA 2020", level=1)
    doc.add_paragraph(PRISMA_ATTRIBUTION).runs[0].italic = True
    _docx_table(doc, [[item["item"], item["topic"], item["status_label"], item["location"] or item["justification"]] for item in data["checklist"]],
                header=["Item", "Tópico", "Situação", "Localização / justificativa"])
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---- LaTeX --------------------------------------------------------------------------------------

_LATEX_SPECIAL = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_", "{": r"\{",
                  "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}


def tex(text: Any) -> str:
    return "".join(_LATEX_SPECIAL.get(char, char) for char in str(text))


def _tex_table(rows: list[list[str]], header: list[str] | None = None) -> str:
    if not rows:
        return "Nenhum registro.\n"
    width = len(header or rows[0])
    lines = [r"\begin{longtable}{" + "|".join(["p{%.2f\\linewidth}" % (0.9 / width)] * width) + "}", r"\hline"]
    if header:
        lines += [" & ".join(r"\textbf{" + tex(cell) + "}" for cell in header) + r" \\", r"\hline"]
    lines += [" & ".join(tex(cell) for cell in row) + r" \\" for row in rows]
    lines += [r"\hline", r"\end{longtable}"]
    return "\n".join(lines) + "\n"


def _tex_flow(flow: dict[str, Any]) -> str:
    sources = r"\\ ".join(f"{tex(item['source'])}: {item['count']}" for item in flow["per_source"])
    return "\n".join([
        r"\begin{center}\begin{tikzpicture}[box/.style={draw, rounded corners, align=center, text width=6cm, minimum height=1cm, font=\small}, node distance=0.8cm and 1cm]",
        rf"\node[box] (id) {{Registros identificados ($n={flow['identified']}$)\\ {sources}}};",
        rf"\node[box, right=of id] (dup) {{Duplicados removidos ($n={flow['duplicates_removed']}$)}};",
        rf"\node[box, below=of id] (scr) {{Registros triados ($n={flow['screened']}$)}};",
        rf"\node[box, right=of scr] (exc) {{Excluídos na seleção ($n={flow['excluded_screening']}$)\\ Aguardando ($n={flow['awaiting_screening']}$)}};",
        rf"\node[box, below=of scr] (qa) {{Avaliados na qualidade ($n={flow['assessed_quality']}$)}};",
        rf"\node[box, right=of qa] (cut) {{Excluídos pela nota de corte ($n={flow['excluded_quality']}$)}};",
        rf"\node[box, below=of qa] (inc) {{Estudos incluídos ($n={flow['included']}$)}};",
        r"\draw[->] (id) -- (scr); \draw[->] (scr) -- (qa); \draw[->] (qa) -- (inc);",
        r"\draw[->] (id) -- (dup); \draw[->] (scr) -- (exc); \draw[->] (qa) -- (cut);",
        r"\end{tikzpicture}\end{center}",
    ]) + "\n"


def _tex_blocks(blocks: list[tuple[str, Any, str]], keys_column: bool = False) -> str:
    out = []
    for title, content, kind in blocks:
        out.append(r"\subsection*{" + tex(title) + "}")
        if kind == "table":
            if keys_column and content and title in {"Estudos incluídos", "Avaliação de qualidade dos incluídos"}:
                rows = [[r"\cite{" + row[0] + "}", *row[1:]] for row in content]
                out.append(_tex_table_raw(rows))
            else:
                out.append(_tex_table(content))
        elif kind == "header_table":
            out.append(_tex_table(content[1:], header=content[0]))
        elif kind == "list":
            out.append("\\begin{itemize}\n" + "\n".join(r"\item " + tex(item) for item in content) + "\n\\end{itemize}\n"
                       if content else "Nenhum registro.\n")
        else:
            out.append("\n\n".join(tex(item) for item in content) + "\n")
    return "\n".join(out)


def _tex_table_raw(rows: list[list[str]]) -> str:
    """Table whose first column already holds LaTeX (a \\cite); the other cells are escaped."""
    width = len(rows[0])
    lines = [r"\begin{longtable}{" + "|".join(["p{%.2f\\linewidth}" % (0.9 / width)] * width) + "}", r"\hline"]
    lines += [" & ".join([row[0], *(tex(cell) for cell in row[1:])]) + r" \\" for row in rows]
    lines += [r"\hline", r"\end{longtable}"]
    return "\n".join(lines) + "\n"


def build_latex(data: dict[str, Any]) -> bytes:
    review, sections, style = data["review"], data["sections"], data["style"]
    preamble = [
        r"\documentclass[11pt,a4paper]{article}",
        r"\usepackage[utf8]{inputenc}", r"\usepackage[T1]{fontenc}", r"\usepackage[brazil]{babel}",
        r"\usepackage{longtable}", r"\usepackage{tikz}", r"\usetikzlibrary{positioning}", r"\usepackage[hidelinks]{hyperref}",
        r"\usepackage[alf]{abntex2cite}" if style == "abnt" else r"\usepackage{natbib}",
        r"\title{" + tex(sections["title"] or review.title) + "}", r"\date{}", r"\begin{document}", r"\maketitle",
    ]
    body = []
    for key, label in REPORT_SECTIONS.items():
        if key == "title":
            continue
        heading = r"\section*{" + tex(label) + "}" if key in {"abstract", "ai_statement"} else r"\section{" + tex(label) + "}"
        body.append(heading)
        body.append(tex(sections[key]) if sections[key] else r"\textit{[A ser escrito pelos pesquisadores.]}")
        if key == "methods":
            body.append(_tex_blocks(methods_blocks(data)))
        if key == "results":
            body.append(r"\subsection*{Fluxograma PRISMA 2020}")
            body.append(_tex_flow(data["flow"]))
            body.append(_tex_blocks(results_blocks(data), keys_column=True))
            if data["articles"]:
                body.append("Estudos incluídos: " + r"\cite{" + ",".join(data["keys"][a.id] for a in data["articles"]) + "}.")
    appendix = [r"\appendix", r"\section{Checklist PRISMA 2020}", r"\textit{" + tex(PRISMA_ATTRIBUTION) + "}",
                _tex_table([[item["item"], item["topic"], item["status_label"], item["location"] or item["justification"]]
                            for item in data["checklist"]], header=["Item", "Tópico", "Situação", "Localização / justificativa"])]
    bib_style = r"\bibliographystyle{abntex2-alf}" if style == "abnt" else r"\bibliographystyle{apalike}"
    main = "\n".join([*preamble, *body, bib_style, r"\bibliography{referencias}", *appendix, r"\end{document}"]) + "\n"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("main.tex", main)
        archive.writestr("referencias.bib", bibtex(data["articles"], data["keys"]))
    return buffer.getvalue()


# ---- export and versions ---------------------------------------------------------------------------

class ExportIn(BaseModel):
    format: Literal["docx", "latex"]


@router.post("/reviews/{review_id}/prisma/export")
def export_report(payload: ExportIn, review: MemberReview, db: DbSession, user: CurrentUser, settings: AppSettings) -> Response:
    data = report_data(db, review)
    content = build_docx(data) if payload.format == "docx" else build_latex(data)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M%S")
    file_name = f"{review.name}-relato-{stamp}.{'docx' if payload.format == 'docx' else 'zip'}"
    object_key = f"reviews/{review.id}/reports/{secrets.token_hex(6)}-{file_name}"
    target = settings.storage_path / object_key
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(content)
    db.add(ReportVersion(review_id=review.id, format=payload.format, reference_style=data["style"], numbers=data["flow"],
                         file_name=file_name, object_key=object_key, created_by_id=user.id))
    log_activity(db, user, f"gerou o relato ({payload.format}) de", review)
    db.commit()
    return Response(content, media_type=DOCX_TYPE if payload.format == "docx" else "application/zip",
                    headers={"Content-Disposition": f'attachment; filename="{file_name}"'})


@router.get("/reviews/{review_id}/prisma/versions")
def list_versions(review: MemberReview, db: DbSession) -> list[dict[str, Any]]:
    return [{"id": item.id, "format": item.format, "reference_style": item.reference_style, "numbers": item.numbers,
             "file_name": item.file_name, "created_by_id": item.created_by_id, "created_at": item.created_at.isoformat()}
            for item in db.scalars(select(ReportVersion).where(ReportVersion.review_id == review.id)
                                   .order_by(ReportVersion.created_at.desc()))]


@router.get("/reviews/{review_id}/prisma/versions/{version_id}/download")
def download_version(version_id: int, review: MemberReview, db: DbSession, settings: AppSettings) -> Response:
    version = db.get(ReportVersion, version_id)
    if version is None or version.review_id != review.id:
        raise HTTPException(status_code=404, detail="Versão não encontrada.")
    path = settings.storage_path / version.object_key
    if not path.is_file():
        raise HTTPException(status_code=410, detail="O arquivo desta versão não está mais disponível.")
    return Response(path.read_bytes(), media_type=DOCX_TYPE if version.format == "docx" else "application/zip",
                    headers={"Content-Disposition": f'attachment; filename="{version.file_name}"'})
